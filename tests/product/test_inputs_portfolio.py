"""Task PF-1: the user-supplied portfolio workbook (`portfolio`) and its validator.

Workbooks here are written by openpyxl, an independent OOXML writer, so the validator is
exercised on bytes it did not produce. Every row is synthetic: a fictional fund and fictional
issuers, no market or fund data.
"""

import io
import shutil
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, get_args
from xml.sax.saxutils import quoteattr
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import openpyxl
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st

from deploy_vii.contracts import Run, Source
from deploy_vii.inputs.portfolio import (
    FIXES,
    MEDIA_TYPE,
    SHEETS,
    STRATEGIES,
    ComplianceRow,
    Mandate,
    PortfolioRefused,
    PortfolioReport,
    portfolio_facts,
    validate_portfolio,
    validate_portfolio_bytes,
)
from deploy_vii.preparation import (
    PreparationRequest,
    SourceRequest,
    load_pack,
    prepare,
    prepare_source,
)
from scripts.build_templates import (
    PORTFOLIO_README,
    PORTFOLIO_ROWS,
    build_portfolio,
    package,
    portfolio_bytes,
    table_sheet,
)

ROOT = Path(__file__).resolve().parents[2]
FILE = "portfolio.xlsx"
TEMPLATE_PATH = "docs/templates/portfolio.template.xlsx"
TEMPLATE = ROOT / TEMPLATE_PATH
# The brief's contract, stated here independently of the module's constants.
BRIEF_SHEETS: dict[str, tuple[str, ...]] = {
    "Mandate": ("section", "field", "value", "notes"),
    "Exposure_Summary": ("metric", "value", "unit", "as_of"),
    "Rating_Distribution": ("rating_bucket", "pct_nav", "issuers"),
    "Sector_Exposure": (
        "sector",
        "pct_nav",
        "issuers",
        "limit_pct_nav",
        "headroom_pct_nav",
        "status",
    ),
    "Top_Holdings": ("rank", "issuer", "instrument", "pct_nav", "rating"),
    "Maturity_Profile": ("bucket", "pct_nav"),
    "Compliance_Monitor": (
        "id",
        "constraint_category",
        "parameter",
        "limit",
        "limit_unit",
        "breach_type",
        "scope",
        "source_document",
        "current_value",
        "headroom",
        "status",
        "last_checked",
        "checked_by",
    ),
    "Holdings_Detail": (
        "issuer",
        "instrument",
        "facility_id",
        "par",
        "market_value",
        "pct_nav",
        "rating",
        "sector",
        "seniority",
        "maturity",
    ),
}
OPTIONAL_COLUMNS = {
    "Mandate": {"notes"},
    "Sector_Exposure": {"limit_pct_nav", "headroom_pct_nav", "status"},
}


def mandate(section: str, field: str, value: object, notes: object = None) -> dict[str, Any]:
    return {"section": section, "field": field, "value": value, "notes": notes}


def control(**values: object) -> dict[str, Any]:
    row = {
        "id": "C-01",
        "constraint_category": "Single Name",
        "parameter": "Max single issuer exposure",
        "limit": "2.5",
        "limit_unit": "% NAV",
        "breach_type": "Hard",
        "scope": "issuer",
        "source_document": "Indenture s7.11 (synthetic)",
        "current_value": "2.37",
        "headroom": "0.13",
        "status": "Watch",
        "last_checked": "2026-05-29",
        "checked_by": "analyst",
    }
    return row | values


def valid() -> dict[str, list[dict[str, Any]]]:
    """A small valid synthetic portfolio, independent of the template builder. Row numbers:
    Mandate nav is C4, nav_currency C5; Exposure_Summary nav is row 2."""
    return {
        "Mandate": [
            mandate("Strategy", "fund_name", "Example Fund I (synthetic)"),
            mandate("Strategy", "strategy", "CLO", "synthetic"),
            mandate("Capital Structure", "nav", 1000),
            mandate("Capital Structure", "nav_currency", "USD"),
            mandate("Reporting", "as_of", "2026-05-29"),
            mandate("Key Parties", "trustee", "Synthetic Trustee Bank"),
        ],
        "Exposure_Summary": [
            {"metric": "nav", "value": "1000.00", "unit": "USD", "as_of": "2026-05-29"},
            {"metric": "obligors", "value": 12, "unit": "count", "as_of": "2026-05-29"},
        ],
        "Rating_Distribution": [
            {"rating_bucket": "B", "pct_nav": "60.5", "issuers": 8},
            {"rating_bucket": "BB", "pct_nav": 39.5, "issuers": "4"},
        ],
        "Sector_Exposure": [
            {
                "sector": "Software",
                "pct_nav": "7.16",
                "issuers": 5,
                "limit_pct_nav": "10.0",
                "headroom_pct_nav": "2.84",
                "status": "Pass",
            },
            {"sector": "Insurance", "pct_nav": "5", "issuers": 2},
        ],
        "Top_Holdings": [
            {
                "rank": 1,
                "issuer": "Alpha Example Holdings (synthetic)",
                "instrument": "Term Loan B due 2030",
                "pct_nav": "2.37",
                "rating": "B3 / B",
            },
            {
                "rank": "2",
                "issuer": "Beta Sample Industries (synthetic)",
                "instrument": "Senior Secured Notes due 2032",
                "pct_nav": "1.91",
                "rating": "Ba3 / BB-",
            },
        ],
        "Maturity_Profile": [
            {"bucket": "3-5 Years", "pct_nav": "40"},
            {"bucket": "5-7 Years", "pct_nav": "60"},
        ],
        "Compliance_Monitor": [
            control(),
            control(
                id="SECTOR",
                constraint_category="Sector",
                parameter="Max single sector",
                limit="10.0",
                breach_type="Soft",
                scope="sector",
                current_value="4.57",
                headroom="5.43",
                status="Pass",
            ),
        ],
    }


HOLDING = {
    "issuer": "Alpha Example Holdings (synthetic)",
    "instrument": "Term Loan B due 2030",
    "facility_id": "TLB-2030",
    "par": "15000000",
    "market_value": "14951250",
    "pct_nav": "2.37",
    "rating": "B3 / B",
    "sector": "Software",
    "seniority": "1st Lien",
    "maturity": "2030-06-30",
}


def workbook(
    data: dict[str, list[dict[str, Any]]],
    *,
    headers: dict[str, tuple[str, ...]] | None = None,
) -> bytes:
    book = openpyxl.Workbook()
    first = book.active
    assert first is not None
    book.remove(first)
    for name, rows in data.items():
        sheet = book.create_sheet(name)
        header = (headers or {}).get(name, BRIEF_SHEETS.get(name, ()))
        if header:
            sheet.append(list(header))
        for row in rows:
            sheet.append([row.get(column) for column in header or BRIEF_SHEETS[name]])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def check(
    data: dict[str, list[dict[str, Any]]], headers: dict[str, tuple[str, ...]] | None = None
) -> PortfolioReport:
    return validate_portfolio_bytes(FILE, workbook(data, headers=headers))


def rendered(report: PortfolioReport) -> list[str]:
    return [issue.render() for issue in report.issues]


def changed(sheet: str, index: int, change: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    data = valid()
    data[sheet][index] = data[sheet][index] | change
    return data


def test_contract_is_the_briefs() -> None:
    assert tuple(sheet.name for sheet in SHEETS) == tuple(BRIEF_SHEETS)
    for sheet in SHEETS:
        assert tuple(column.name for column in sheet.columns) == BRIEF_SHEETS[sheet.name]
        optional = {column.name for column in sheet.columns if not column.required}
        assert optional == OPTIONAL_COLUMNS.get(sheet.name, set()), sheet.name
    assert [sheet.name for sheet in SHEETS if sheet.optional] == ["Holdings_Detail"]
    assert MEDIA_TYPE == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_typed_row_literals_match_the_validators_choices() -> None:
    # The typed rows' Literal fields must accept exactly what the validator passes.
    fields = ComplianceRow.model_fields
    choices = {column.name: column.choices for column in SHEETS[6].columns}
    assert get_args(fields["scope"].annotation) == choices["scope"]
    assert get_args(fields["breach_type"].annotation) == choices["breach_type"]
    assert get_args(fields["status"].annotation) == choices["status"]
    assert get_args(Mandate.model_fields["strategy"].annotation) == STRATEGIES


def test_comments_and_merged_ranges_are_not_cells() -> None:
    book = openpyxl.load_workbook(io.BytesIO(workbook(valid())))
    sheet = book["Mandate"]
    sheet["D2"].comment = openpyxl.comments.Comment("synthetic note", "analyst")
    sheet.merge_cells("D3:D4")
    stream = io.BytesIO()
    book.save(stream)
    raw = stream.getvalue()
    spans = [span.locator for span in source(raw).spans]
    assert any(locator.startswith("part:") for locator in spans)
    assert "sheet:Mandate!merged:D3:D4" in spans
    assert rendered(validate_portfolio_bytes(FILE, raw)) == []


def test_a_valid_workbook_is_ok_with_sheet_and_compliance_counts() -> None:
    report = check(valid())
    assert rendered(report) == []
    assert report.ok
    assert (report.sheets, report.compliance_rows, report.rows) == (7, 2, 18)
    assert report.lines() == ("OK: 7 sheets, 2 compliance rows",)


def test_the_optional_holdings_sheet_is_counted_when_present() -> None:
    report = check(valid() | {"Holdings_Detail": [HOLDING]})
    assert rendered(report) == []
    assert report.lines() == ("OK: 8 sheets, 2 compliance rows",)


NUMBER_FIX = (
    "enter a plain number: digits, an optional sign and one decimal point "
    "(no %, units, spaces or thousands separators)"
)
DATE_FIX = "enter the date as text YYYY-MM-DD (format the column as Text so Excel keeps it)"
PCT_FIX = "enter the share of NAV in percent as a plain number (2.5 means 2.5 % NAV)"
COUNT_FIX = "enter the number of issuers as a whole number of at least 0, e.g. 12"
STATUS_FIX = "enter Pass, Watch or Breach"
FORMULA_FIX = "replace the formula with its value (Paste Special, Values)"
SECTIONS = (
    "Strategy, Key Parties, Capital Structure, Mandate Constraints, Liquidity Terms, Fees "
    "or Reporting"
)
SCOPES = (
    "issuer, rating_bucket, sector, seniority, instrument, duration, diversification, "
    "yield, pricing, covenant, esg or other"
)


@pytest.mark.parametrize(
    "sheet, index, change, cell, problem, fix",
    [
        (
            "Mandate",
            5,
            {"section": "Parties"},
            "A7",
            f"section 'Parties' is not one of {SECTIONS}",
            f"enter one of {SECTIONS}",
        ),
        ("Mandate", 5, {"section": None}, "A7", "section is empty", f"enter one of {SECTIONS}"),
        (
            "Mandate",
            5,
            {"field": None},
            "B7",
            "field is empty",
            "enter the field's name, e.g. fund_name",
        ),
        (
            "Mandate",
            5,
            {"value": " "},
            "C7",
            "value is empty",
            "enter the field's value, or delete the row",
        ),
        (
            "Mandate",
            1,
            {"value": "Loans"},
            "C3",
            "strategy 'Loans' is not one of CLO, Indexed Loans, Indexed HY or Other",
            "enter one of CLO, Indexed Loans, Indexed HY or Other",
        ),
        ("Mandate", 2, {"value": "1,000"}, "C4", "nav '1,000' is not a number", NUMBER_FIX),
        (
            "Mandate",
            2,
            {"value": "0"},
            "C4",
            "nav 0 is not above 0",
            "enter the fund's NAV, a number above 0",
        ),
        (
            "Mandate",
            3,
            {"value": "US$"},
            "C5",
            "nav_currency 'US$' is not a three-letter ISO 4217 code",
            "enter the three-letter upper-case ISO 4217 code, e.g. USD",
        ),
        (
            "Mandate",
            4,
            {"value": "29-May-2026"},
            "C6",
            "as_of '29-May-2026' is not an ISO date YYYY-MM-DD",
            DATE_FIX,
        ),
        (
            "Mandate",
            4,
            {"value": 46171},
            "C6",
            "as_of '46171' is an Excel date serial, not an ISO date YYYY-MM-DD",
            DATE_FIX,
        ),
        (
            "Exposure_Summary",
            1,
            {"metric": None},
            "A3",
            "metric is empty",
            "enter the metric's name, e.g. nav",
        ),
        (
            "Exposure_Summary",
            1,
            {"value": "350 obligors"},
            "B3",
            "value '350 obligors' is not a number",
            NUMBER_FIX,
        ),
        (
            "Exposure_Summary",
            1,
            {"value": None},
            "B3",
            "value is empty",
            "enter the metric's value as a plain number",
        ),
        (
            "Exposure_Summary",
            1,
            {"unit": None},
            "C3",
            "unit is empty",
            "enter the unit, e.g. USD, % NAV, years or count",
        ),
        (
            "Exposure_Summary",
            1,
            {"as_of": "2026-13-01"},
            "D3",
            "as_of '2026-13-01' is not an ISO date YYYY-MM-DD",
            DATE_FIX,
        ),
        (
            "Rating_Distribution",
            0,
            {"rating_bucket": None},
            "A2",
            "rating_bucket is empty",
            "enter the rating bucket, e.g. B or Unrated",
        ),
        ("Rating_Distribution", 0, {"pct_nav": None}, "B2", "pct_nav is empty", PCT_FIX),
        (
            "Rating_Distribution",
            0,
            {"issuers": "-1"},
            "C2",
            "issuers '-1' is not a whole number of at least 0",
            COUNT_FIX,
        ),
        (
            "Rating_Distribution",
            0,
            {"issuers": "8.5"},
            "C2",
            "issuers '8.5' is not a whole number of at least 0",
            COUNT_FIX,
        ),
        (
            "Sector_Exposure",
            0,
            {"sector": None},
            "A2",
            "sector is empty",
            "enter the sector's name",
        ),
        (
            "Sector_Exposure",
            0,
            {"pct_nav": "7.16%"},
            "B2",
            "pct_nav '7.16%' is not a number",
            NUMBER_FIX,
        ),
        ("Sector_Exposure", 1, {"issuers": None}, "C3", "issuers is empty", COUNT_FIX),
        (
            "Sector_Exposure",
            0,
            {"limit_pct_nav": "≤ 10.0%"},
            "D2",
            "limit_pct_nav '≤ 10.0%' is not a number",
            NUMBER_FIX,
        ),
        (
            "Sector_Exposure",
            0,
            {"status": "OK"},
            "F2",
            "status 'OK' is not Pass, Watch or Breach",
            STATUS_FIX,
        ),
        (
            "Top_Holdings",
            0,
            {"rank": "0"},
            "A2",
            "rank '0' is not a whole number of at least 1",
            "number the holdings 1, 2, 3 and so on, with no gaps or repeats",
        ),
        ("Top_Holdings", 0, {"issuer": None}, "B2", "issuer is empty", "enter the issuer's name"),
        (
            "Top_Holdings",
            0,
            {"instrument": None},
            "C2",
            "instrument is empty",
            "describe the instrument, e.g. Term Loan B due 2030",
        ),
        ("Top_Holdings", 0, {"pct_nav": "n/a"}, "D2", "pct_nav 'n/a' is not a number", NUMBER_FIX),
        (
            "Top_Holdings",
            0,
            {"rating": None},
            "E2",
            "rating is empty",
            "enter the rating as reported, e.g. B2 / B",
        ),
        (
            "Maturity_Profile",
            0,
            {"bucket": None},
            "A2",
            "bucket is empty",
            "enter the maturity bucket, e.g. 3-5 Years",
        ),
        (
            "Maturity_Profile",
            1,
            {"pct_nav": "1e5000"},
            "B3",
            "pct_nav '1e5000' is not a number",
            NUMBER_FIX,
        ),
        (
            "Compliance_Monitor",
            0,
            {"id": "C1"},
            "A2",
            "id 'C1' is not of the form C-01 or SECTOR",
            "enter an id such as C-01 (capital letters, a hyphen, two digits), or SECTOR for "
            "the sector-limit row",
        ),
        (
            "Compliance_Monitor",
            0,
            {"constraint_category": None},
            "B2",
            "constraint_category is empty",
            "enter the constraint category, e.g. Single Name",
        ),
        (
            "Compliance_Monitor",
            0,
            {"parameter": None},
            "C2",
            "parameter is empty",
            "describe the constraint, e.g. Max single issuer exposure",
        ),
        (
            "Compliance_Monitor",
            0,
            {"limit": "≤ 2.5% NAV"},
            "D2",
            "limit '≤ 2.5% NAV' is not a number",
            NUMBER_FIX,
        ),
        (
            "Compliance_Monitor",
            0,
            {"limit_unit": "%NAV"},
            "E2",
            "limit_unit '%NAV' is not % NAV or a three-letter currency code",
            "enter % NAV, or the Mandate nav_currency code (e.g. USD) for an amount limit",
        ),
        (
            "Compliance_Monitor",
            0,
            {"breach_type": "—"},
            "F2",
            "breach_type '—' is not Hard or Soft",
            "enter Hard or Soft",
        ),
        (
            "Compliance_Monitor",
            0,
            {"scope": "Issuer"},
            "G2",
            f"scope 'Issuer' is not one of {SCOPES}",
            f"enter one of {SCOPES}",
        ),
        (
            "Compliance_Monitor",
            0,
            {"source_document": None},
            "H2",
            "source_document is empty",
            "enter the document that sets the limit, e.g. Indenture s7.11",
        ),
        (
            "Compliance_Monitor",
            0,
            {"current_value": None},
            "I2",
            "current_value is empty",
            "enter the current exposure as a plain number in limit_unit",
        ),
        (
            "Compliance_Monitor",
            0,
            {"headroom": "0.13%"},
            "J2",
            "headroom '0.13%' is not a number",
            NUMBER_FIX,
        ),
        (
            "Compliance_Monitor",
            0,
            {"status": None},
            "K2",
            "status is empty",
            STATUS_FIX,
        ),
        (
            "Compliance_Monitor",
            0,
            {"last_checked": "29-May-2026"},
            "L2",
            "last_checked '29-May-2026' is not an ISO date YYYY-MM-DD",
            DATE_FIX,
        ),
        (
            "Compliance_Monitor",
            0,
            {"checked_by": None},
            "M2",
            "checked_by is empty",
            "enter who checked the row, e.g. analyst",
        ),
        (
            "Compliance_Monitor",
            0,
            {"limit": "=2+0.5"},
            "D2",
            "limit holds a formula; prepare freezes stored values and never recalculates",
            FORMULA_FIX,
        ),
    ],
)
def test_each_column_rule_fails_with_its_exact_message(
    sheet: str, index: int, change: dict[str, Any], cell: str, problem: str, fix: str
) -> None:
    report = check(changed(sheet, index, change))
    assert rendered(report) == [f"{FILE}:{sheet}!{cell}: {problem} — fix: {fix}"]
    assert report.lines()[-1] == "FAILED: 1 issues"


@pytest.mark.parametrize(
    "column, change, cell, problem, fix",
    [
        ("issuer", None, "A2", "issuer is empty", "enter the issuer's name"),
        (
            "facility_id",
            "TLB 2030",
            "C2",
            "facility_id 'TLB 2030' has characters other than letters, digits, - and _",
            "enter the facility's id using letters, digits, - and _ only, e.g. TLB-2030",
        ),
        ("par", "15,000,000", "D2", "par '15,000,000' is not a number", NUMBER_FIX),
        (
            "market_value",
            None,
            "E2",
            "market_value is empty",
            "enter the market value as a plain number",
        ),
        (
            "seniority",
            None,
            "I2",
            "seniority is empty",
            "enter the seniority, e.g. 1st Lien",
        ),
        (
            "maturity",
            47664,
            "J2",
            "maturity '47664' is an Excel date serial, not an ISO date YYYY-MM-DD",
            DATE_FIX,
        ),
    ],
)
def test_holdings_detail_column_rules(
    column: str, change: object, cell: str, problem: str, fix: str
) -> None:
    report = check(valid() | {"Holdings_Detail": [HOLDING | {column: change}]})
    assert rendered(report) == [f"{FILE}:Holdings_Detail!{cell}: {problem} — fix: {fix}"]


@pytest.mark.parametrize(
    "field, row, fix",
    [
        ("fund_name", 0, "add a row: section Strategy, field fund_name, value the fund's name"),
        (
            "strategy",
            1,
            "add a row: section Strategy, field strategy, value CLO, Indexed Loans, Indexed HY "
            "or Other",
        ),
        (
            "nav",
            2,
            "add a row: section Capital Structure, field nav, value the fund's NAV as a plain "
            "number",
        ),
        (
            "nav_currency",
            3,
            "add a row: section Capital Structure, field nav_currency, value the NAV's "
            "three-letter currency code, e.g. USD",
        ),
        (
            "as_of",
            4,
            "add a row: section Reporting, field as_of, value the NAV and compliance date as "
            "text YYYY-MM-DD",
        ),
    ],
)
def test_a_missing_required_mandate_field_is_one_issue(field: str, row: int, fix: str) -> None:
    data = valid()
    section = data["Mandate"][row]["section"]
    del data["Mandate"][row]
    report = check(data)
    assert rendered(report) == [
        f"{FILE}:Mandate!A7: no {section} row with field {field} — fix: {fix}"
    ]


def test_a_required_field_in_the_wrong_section_is_missing() -> None:
    report = check(changed("Mandate", 4, {"section": "Strategy"}))
    assert [issue.problem for issue in report.issues] == ["no Reporting row with field as_of"]


def test_a_repeated_mandate_field_is_one_issue_on_the_later_row() -> None:
    data = valid()
    data["Mandate"].append(mandate("Capital Structure", "nav", 1000))
    assert rendered(check(data)) == [
        f"{FILE}:Mandate!A8:B8: duplicates row 4 on section and field — fix: keep one row per "
        "section and field, or correct the repeated value"
    ]


def test_exposure_nav_must_equal_the_mandate_nav() -> None:
    assert rendered(check(changed("Exposure_Summary", 0, {"value": "1000.01"}))) == [
        f"{FILE}:Exposure_Summary!B2: nav 1000.01 is not the Mandate nav 1000 (Mandate!C4)"
        " — fix: enter the same NAV on Exposure_Summary and Mandate"
    ]
    # Same number, different lexeme: equal as decimals.
    assert rendered(check(changed("Exposure_Summary", 0, {"value": "1e3"}))) == []


def test_exposure_nav_unit_must_be_the_mandate_currency() -> None:
    assert rendered(check(changed("Exposure_Summary", 0, {"unit": "EUR"}))) == [
        f"{FILE}:Exposure_Summary!C2: unit 'EUR' of nav is not the Mandate nav_currency USD"
        " — fix: enter the nav row's unit as the Mandate nav_currency, e.g. USD"
    ]


def test_exposure_without_nav_is_one_issue() -> None:
    data = valid()
    del data["Exposure_Summary"][0]
    assert rendered(check(data)) == [
        f"{FILE}:Exposure_Summary!A3: no metric nav — fix: add a row: metric nav, value the "
        "fund's NAV (the Mandate nav), unit the nav_currency, as_of the NAV date"
    ]


def test_one_pct_nav_is_computed_never_supplied() -> None:
    data = valid()
    data["Exposure_Summary"].append(
        {"metric": "one_pct_nav", "value": "10", "unit": "USD", "as_of": "2026-05-29"}
    )
    assert rendered(check(data)) == [
        f"{FILE}:Exposure_Summary!A4: one_pct_nav is computed from nav and is not supplied"
        " — fix: delete the one_pct_nav row; Python computes 1 % of NAV from nav"
    ]


def test_a_repeated_metric_is_one_issue() -> None:
    data = valid()
    data["Exposure_Summary"].append(data["Exposure_Summary"][1])
    assert [issue.render() for issue in check(data).issues] == [
        f"{FILE}:Exposure_Summary!A4: duplicates row 3 on metric — fix: keep one row per "
        "metric, or correct the repeated value"
    ]


@pytest.mark.parametrize(
    "rows, ok",
    [
        ([("B", "60"), ("BB", "39")], True),
        ([("B", "60"), ("BB", "41")], True),
        ([("B", "60"), ("BB", "38.99")], False),
        ([("B", "60"), ("BB", "41.01")], False),
        ([("B", "60"), ("BB", "30"), ("Unallocated", "10")], True),
    ],
)
def test_rating_distribution_sums_to_between_99_and_101(
    rows: list[tuple[str, str]], ok: bool
) -> None:
    data = valid()
    data["Rating_Distribution"] = [
        {"rating_bucket": bucket, "pct_nav": pct, "issuers": 1} for bucket, pct in rows
    ]
    report = check(data)
    if ok:
        assert rendered(report) == []
    else:
        total = sum((Decimal(pct) for _, pct in rows), Decimal(0))
        assert rendered(report) == [
            f"{FILE}:Rating_Distribution!B2:B3: pct_nav sums to {total}, outside 99.0 to 101.0"
            " — fix: correct the percentages, or add a row Unallocated with the remainder of NAV"
        ]


@pytest.mark.parametrize(
    "sheet, column, label, cell",
    [
        ("Rating_Distribution", "rating_bucket", "Total", "A2"),
        ("Sector_Exposure", "sector", "Top 20 Subtotal", "A2"),
        ("Maturity_Profile", "bucket", "total", "A2"),
    ],
)
def test_total_rows_are_refused(sheet: str, column: str, label: str, cell: str) -> None:
    report = check(changed(sheet, 0, {column: label}))
    assert rendered(report) == [
        f"{FILE}:{sheet}!{cell}: {column} '{label}' is a total row — fix: leave out total and "
        "subtotal rows; Python adds the rows up"
    ]


@pytest.mark.parametrize(
    "sheet, column, letter",
    [
        ("Rating_Distribution", "rating_bucket", "A"),
        ("Sector_Exposure", "sector", "A"),
        ("Maturity_Profile", "bucket", "A"),
        ("Top_Holdings", "rank", "A"),
    ],
)
def test_repeated_keys_are_one_issue(sheet: str, column: str, letter: str) -> None:
    data = valid()
    data[sheet][1] = data[sheet][1] | {column: data[sheet][0][column]}
    if sheet == "Rating_Distribution":
        data[sheet][1]["pct_nav"] = "39.5"
    assert rendered(check(data)) == [
        f"{FILE}:{sheet}!{letter}3: duplicates row 2 on {column} — fix: keep one row per "
        f"{column}, or correct the repeated value"
    ]


def test_ranks_must_run_from_one_without_gaps() -> None:
    report = check(changed("Top_Holdings", 1, {"rank": 3}))
    assert rendered(report) == [
        f"{FILE}:Top_Holdings!A2:A3: ranks are not 1 to 2: 2 missing — fix: number the "
        "holdings 1, 2, 3 and so on, with no gaps or repeats"
    ]


def test_holdings_detail_repeated_facility_is_one_issue() -> None:
    report = check(valid() | {"Holdings_Detail": [HOLDING, HOLDING]})
    assert rendered(report) == [
        f"{FILE}:Holdings_Detail!C3: duplicates row 2 on facility_id — fix: keep one row per "
        "facility_id, or correct the repeated value"
    ]


COMPLIANCE_HEADROOM_FIX = "enter limit minus current_value, or correct limit or current_value"


def test_compliance_headroom_must_be_limit_minus_current() -> None:
    assert rendered(check(changed("Compliance_Monitor", 0, {"headroom": "0.15"}))) == [
        f"{FILE}:Compliance_Monitor!J2: headroom 0.15 is not limit 2.5 minus current_value 2.37"
        f" = 0.13 (to within 0.01) — fix: {COMPLIANCE_HEADROOM_FIX}"
    ]
    # Within 0.01 either way is accepted.
    assert rendered(check(changed("Compliance_Monitor", 0, {"headroom": "0.14"}))) == []
    assert rendered(check(changed("Compliance_Monitor", 0, {"headroom": "0.12"}))) == []


def test_currency_limits_need_the_mandate_currency_and_skip_headroom_arithmetic() -> None:
    row = {
        "limit": "20000000",
        "limit_unit": "USD",
        "current_value": "15000000",
        "headroom": "1",
        "status": "Pass",
    }
    assert rendered(check(changed("Compliance_Monitor", 0, row))) == []
    assert rendered(check(changed("Compliance_Monitor", 0, row | {"limit_unit": "EUR"}))) == [
        f"{FILE}:Compliance_Monitor!E2: limit_unit EUR is not the Mandate nav_currency USD"
        " — fix: enter % NAV, or the Mandate nav_currency code (e.g. USD) for an amount limit"
    ]


@pytest.mark.parametrize(
    "change, problem, fix",
    [
        (
            {"current_value": "2.6", "headroom": "-0.1", "status": "Watch"},
            "status 'Watch' but current_value 2.6 exceeds limit 2.5",
            "enter Breach, or correct the values: a current value above its limit is a breach",
        ),
        (
            {"status": "Breach"},
            "status 'Breach' but current_value 2.37 does not exceed limit 2.5",
            "enter Pass or Watch, or correct the values: only a current value above its limit "
            "is a breach",
        ),
        (
            {"current_value": "2.25", "headroom": "0.25", "status": "Watch"},
            "status 'Watch' but the headroom 0.25 is at least 10 % of limit 2.5",
            "enter Pass: Watch is for headroom under 10 % of the limit",
        ),
    ],
)
def test_compliance_status_must_follow_the_values(
    change: dict[str, Any], problem: str, fix: str
) -> None:
    assert rendered(check(changed("Compliance_Monitor", 0, change))) == [
        f"{FILE}:Compliance_Monitor!K2: {problem} — fix: {fix}"
    ]


@pytest.mark.parametrize(
    "change",
    [
        {"current_value": "2.6", "headroom": "-0.1", "status": "Breach"},
        {"current_value": "2.5", "headroom": "0", "status": "Watch"},
        {"current_value": "2.26", "headroom": "0.24", "status": "Watch"},
        {"current_value": "2.37", "headroom": "0.13", "status": "Pass"},
    ],
)
def test_compliance_status_accepted_combinations(change: dict[str, Any]) -> None:
    assert rendered(check(changed("Compliance_Monitor", 0, change))) == []


def test_sector_headroom_and_status_rules() -> None:
    def sector(change: dict[str, Any]) -> list[str]:
        return [issue.render() for issue in check(changed("Sector_Exposure", 0, change)).issues]

    assert sector({"headroom_pct_nav": "2.9"}) == [
        f"{FILE}:Sector_Exposure!E2: headroom_pct_nav 2.9 is not limit_pct_nav 10.0 minus "
        "pct_nav 7.16 = 2.84 (to within 0.01) — fix: enter limit_pct_nav minus pct_nav, or "
        "correct the limit or the exposure"
    ]
    assert sector({"limit_pct_nav": None, "status": None}) == [
        f"{FILE}:Sector_Exposure!E2: headroom_pct_nav is given without limit_pct_nav — fix: "
        "enter limit_pct_nav, or leave headroom_pct_nav empty"
    ]
    assert sector({"limit_pct_nav": None, "headroom_pct_nav": None}) == [
        f"{FILE}:Sector_Exposure!F2: status is given without limit_pct_nav — fix: enter "
        "limit_pct_nav, or leave status empty"
    ]
    assert sector({"status": "Breach"}) == [
        f"{FILE}:Sector_Exposure!F2: status 'Breach' but pct_nav 7.16 does not exceed "
        "limit_pct_nav 10.0 — fix: enter Pass or Watch, or correct the values: only a current "
        "value above its limit is a breach"
    ]
    assert sector({"headroom_pct_nav": None, "status": None}) == []


def test_numbers_are_exact_decimals_never_floats() -> None:
    # In binary floating point 0.3 - 0.1 - 0.19 is not 0.01; as decimals the headroom 0.02
    # differs from 0.3 - 0.29 = 0.01 by exactly 0.01, inside the tolerance.
    change = {"limit": "0.3", "current_value": "0.29", "headroom": "0.02", "status": "Watch"}
    assert rendered(check(changed("Compliance_Monitor", 0, change))) == []
    change |= {"headroom": "0.0200000000000000000001"}
    assert [issue.cell for issue in check(changed("Compliance_Monitor", 0, change)).issues] == [
        "J2"
    ]


def test_a_number_longer_than_40_characters_is_refused() -> None:
    long = "1" + "0" * 40
    assert rendered(check(changed("Exposure_Summary", 1, {"value": long}))) == [
        f"{FILE}:Exposure_Summary!B3: value '{long}' is longer than 40 characters — fix: enter "
        "the number with at most 40 characters (Excel keeps 15 significant digits)"
    ]


def source(raw: bytes, name: str = FILE) -> Source:
    request = SourceRequest(source_id="portfolio", path=name, media_type=MEDIA_TYPE)
    return prepare_source(request, raw)


def test_portfolio_facts_returns_typed_decimal_rows_with_locators() -> None:
    facts = portfolio_facts(source(workbook(valid() | {"Holdings_Detail": [HOLDING]})))
    assert facts.mandate.fund_name == "Example Fund I (synthetic)"
    assert facts.mandate.strategy == "CLO"
    assert facts.mandate.nav == Decimal("1000")
    assert facts.mandate.nav_currency == "USD"
    assert facts.mandate.as_of == date(2026, 5, 29)
    assert [(item.section, item.field) for item in facts.mandate.fields][-1] == (
        "Key Parties",
        "trustee",
    )
    assert facts.mandate.locator("nav") == "sheet:Mandate!C4"
    nav = facts.exposures[0]
    assert (nav.metric, nav.value, nav.unit, nav.as_of) == (
        "nav",
        Decimal("1000.00"),
        "USD",
        date(2026, 5, 29),
    )
    assert nav.locator("value") == "sheet:Exposure_Summary!B2"
    assert [(item.rating_bucket, item.pct_nav, item.issuers) for item in facts.ratings] == [
        ("B", Decimal("60.5"), 8),
        ("BB", Decimal("39.5"), 4),
    ]
    software, insurance = facts.sectors
    assert (software.limit_pct_nav, software.headroom_pct_nav, software.status) == (
        Decimal("10.0"),
        Decimal("2.84"),
        "Pass",
    )
    assert (insurance.limit_pct_nav, insurance.headroom_pct_nav, insurance.status) == (
        None,
        None,
        None,
    )
    assert [item.rank for item in facts.top_holdings] == [1, 2]
    assert [item.pct_nav for item in facts.maturities] == [Decimal("40"), Decimal("60")]
    issuer, sector = facts.compliance
    assert (issuer.id, issuer.limit, issuer.current_value, issuer.headroom) == (
        "C-01",
        Decimal("2.5"),
        Decimal("2.37"),
        Decimal("0.13"),
    )
    assert (issuer.breach_type, issuer.scope, issuer.status, issuer.last_checked) == (
        "Hard",
        "issuer",
        "Watch",
        date(2026, 5, 29),
    )
    assert sector.locator("current_value") == "sheet:Compliance_Monitor!I3"
    assert facts.holdings is not None
    (holding,) = facts.holdings
    assert (holding.par, holding.market_value, holding.maturity) == (
        Decimal("15000000"),
        Decimal("14951250"),
        date(2030, 6, 30),
    )
    every = [
        value
        for row in (*facts.exposures, *facts.ratings, *facts.sectors, *facts.compliance)
        for value in row.model_dump().values()
    ]
    assert not any(isinstance(value, float) for value in every)


def test_portfolio_facts_without_the_optional_sheet_has_no_holdings() -> None:
    assert portfolio_facts(source(workbook(valid()))).holdings is None


def test_portfolio_facts_refuses_an_invalid_workbook_with_its_report() -> None:
    raw = workbook(changed("Compliance_Monitor", 0, {"headroom": "0.15"}))
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(source(raw))
    assert refused.value.report == validate_portfolio_bytes(FILE, raw)
    assert [issue.cell for issue in refused.value.report.issues] == ["J2"]


def test_validate_portfolio_reads_the_file(tmp_path: Path) -> None:
    path = tmp_path / FILE
    path.write_bytes(workbook(valid()))
    assert validate_portfolio(path) == validate_portfolio_bytes(FILE, path.read_bytes())
    missing = validate_portfolio(tmp_path / "absent.xlsx")
    assert rendered(missing) == [
        "absent.xlsx:Mandate!A1: cannot read the file: No such file or directory"
        " — fix: check the path and that the file is readable"
    ]


def test_fixes_are_unique_strings() -> None:
    assert len(set(FIXES)) == len(FIXES)
    assert TEMPLATE_PATH in " ".join(FIXES)


def run_cli(arguments: list[str], monkeypatch: pytest.MonkeyPatch) -> int:
    from deploy_vii.__main__ import main

    monkeypatch.setattr(sys, "argv", ["deploy-vii", "validate-input", *arguments])
    with pytest.raises(SystemExit) as exit_info:
        main()
    code = exit_info.value.code
    assert isinstance(code, int)
    return code


def test_cli_portfolio_prints_ok_and_exits_zero(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / FILE
    path.write_bytes(workbook(valid()))
    assert run_cli(["portfolio", str(path)], monkeypatch) == 0
    assert capsys.readouterr().out == "OK: 7 sheets, 2 compliance rows\n"


def test_cli_portfolio_prints_one_line_per_issue_and_exits_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    data = changed("Compliance_Monitor", 0, {"headroom": "0.15", "checked_by": None})
    path = tmp_path / FILE
    path.write_bytes(workbook(data))
    assert run_cli(["portfolio", str(path)], monkeypatch) == 1
    assert capsys.readouterr().out.splitlines() == [
        f"{FILE}:Compliance_Monitor!M2: checked_by is empty — fix: enter who checked the row,"
        " e.g. analyst",
        f"{FILE}:Compliance_Monitor!J2: headroom 0.15 is not limit 2.5 minus current_value 2.37"
        f" = 0.13 (to within 0.01) — fix: {COMPLIANCE_HEADROOM_FIX}",
        "FAILED: 2 issues",
    ]


def test_cli_portfolio_json_prints_the_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / FILE
    path.write_bytes(workbook(changed("Exposure_Summary", 0, {"value": "999"})))
    assert run_cli(["portfolio", str(path), "--json"], monkeypatch) == 1
    printed = PortfolioReport.model_validate_json(capsys.readouterr().out)
    assert printed == validate_portfolio(path)
    assert (printed.ok, printed.sheets, printed.compliance_rows) == (False, 7, 2)


def test_cli_portfolio_kind_comes_from_the_input_table() -> None:
    from deploy_vii.inputs import KINDS

    assert tuple(KINDS) == ("rv", "portfolio")
    assert KINDS["portfolio"] is validate_portfolio


def test_template_validates_ok_with_eight_sheets_and_the_design_limits() -> None:
    report = validate_portfolio(TEMPLATE)
    assert rendered(report) == []
    assert report.lines() == ("OK: 8 sheets, 5 compliance rows",)
    facts = portfolio_facts(source(TEMPLATE.read_bytes(), TEMPLATE.name))
    assert facts.mandate.nav == Decimal("630928502")
    assert facts.mandate.nav_currency == "USD"
    assert facts.mandate.as_of == date(2026, 5, 29)
    limits = {
        row.id: (row.limit, row.limit_unit, row.breach_type, row.scope, row.current_value)
        for row in facts.compliance
    }
    assert limits == {
        "C-01": (Decimal("2.5"), "% NAV", "Hard", "issuer", Decimal("2.37")),
        "C-04": (Decimal("7.5"), "% NAV", "Hard", "rating_bucket", Decimal("7.36")),
        "C-05": (Decimal("5.0"), "% NAV", "Soft", "rating_bucket", Decimal("3.73")),
        "C-10": (Decimal("10.0"), "% NAV", "Soft", "seniority", Decimal("0.61")),
        "SECTOR": (Decimal("10.0"), "% NAV", "Soft", "sector", Decimal("4.57")),
    }
    assert {row.checked_by for row in facts.compliance} == {"analyst"}
    assert all(row.last_checked == date(2026, 5, 29) for row in facts.compliance)
    issuers = [row.issuer for row in (*facts.top_holdings, *(facts.holdings or ()))]
    assert issuers and all(issuer.endswith("(synthetic)") for issuer in issuers)
    assert facts.mandate.fund_name.endswith("(synthetic)")


def test_builder_is_byte_stable_and_matches_the_committed_template(tmp_path: Path) -> None:
    build_portfolio(tmp_path / "first.xlsx")
    build_portfolio(tmp_path / "second.xlsx")
    first = (tmp_path / "first.xlsx").read_bytes()
    assert first == (tmp_path / "second.xlsx").read_bytes()
    assert first == portfolio_bytes(PORTFOLIO_ROWS)
    assert first == TEMPLATE.read_bytes(), "run: uv run python scripts/build_templates.py"


def test_an_independent_reader_sees_the_template_sheets_headers_and_readme() -> None:
    book = openpyxl.load_workbook(TEMPLATE, read_only=True)
    assert book.sheetnames == [*BRIEF_SHEETS, "README"]
    for name, header in BRIEF_SHEETS.items():
        rows = list(book[name].iter_rows(values_only=True))
        assert rows[0] == header, name
        assert len(rows) == 1 + len(PORTFOLIO_ROWS[name]), name
    readme = [row[0] for row in book["README"].iter_rows(values_only=True)]
    assert readme == list(PORTFOLIO_README)
    assert readme[0] == "synthetic example — replace every row"
    compliance = list(book["Compliance_Monitor"].iter_rows(values_only=True))
    assert compliance[1][3] == 2.5  # limits are numeric cells; dates and ids are text
    assert compliance[1][11] == "2026-05-29"


# Every text cell the template may hold besides its headers and README lines, written out here
# so that any new string in the builder needs a reviewed edit of this list. Fictional names
# only; none is a reference workbook's person, party or issuer.
TEMPLATE_TEXT = (
    # Mandate
    "Strategy",
    "Key Parties",
    "Capital Structure",
    "Mandate Constraints",
    "Liquidity Terms",
    "Fees",
    "Reporting",
    "fund_name",
    "strategy",
    "non_call_period",
    "legal_final_maturity",
    "collateral_manager",
    "trustee",
    "nav",
    "nav_currency",
    "subordinated_notes",
    "constraint_schedule",
    "redemption",
    "senior_management_fee",
    "as_of",
    "investor_reporting",
    "Test CLO I Ltd (synthetic)",
    "CLO",
    "2025-04-15 to 2027-04-15",
    "2038-04-15",
    "Synthetic Asset Management LLP",
    "Synthetic Trustee Bank",
    "630928502",
    "USD",
    "USD 45,000,000",
    "see Compliance_Monitor",
    "none (CLO structure)",
    "20 bps p.a. on collateral balance",
    "2026-05-29",
    "monthly trustee report",
    "fictional fund",
    "2 years",
    "collateral pool market value",
    "first loss",
    "callable after NC",
    "NAV and compliance date",
    # Exposure_Summary
    "total_par",
    "obligors",
    "positions",
    "sectors",
    "war_notch",
    "wa_margin",
    "wa_ytm",
    "wal",
    "wa_price",
    "first_lien_pct_nav",
    "top10_pct_nav",
    "count",
    "rating notch",
    "bps",
    "%",
    "years",
    "price",
    "% NAV",
    # Rating_Distribution, Sector_Exposure, Maturity_Profile
    "IG (Baa+)",
    "BB (Ba)",
    "B",
    "CCC (Caa+)",
    "Other",
    "Unrated",
    "Health Care Services",
    "Business Services",
    "Software",
    "Chemicals",
    "Leisure Facilities & Services",
    "Pass",
    "< 1 Year",
    "1-2 Years",
    "2-3 Years",
    "3-5 Years",
    "5-7 Years",
    # Top_Holdings and Holdings_Detail
    "Alpha Example Holdings (synthetic)",
    "Beta Sample Aerospace (synthetic)",
    "Gamma Illustrative Media (synthetic)",
    "Delta Placeholder Networks (synthetic)",
    "Epsilon Demo Leisure (synthetic)",
    "Term Loan B due 2030 (synthetic)",
    "Term Loan B due 2031 (synthetic)",
    "Term Loan B due 2029 (synthetic)",
    "Term Loan due 2028 (synthetic)",
    "Senior Secured Notes due 2032 (synthetic)",
    "B2 / B+",
    "Ba2 / BB",
    "B3 / B-",
    "B1 / B",
    "Ba3 / BB-",
    "TLB-2030",
    "TLB-2031",
    "TLB-2029",
    "TL-2028",
    "SSN-2032",
    "1st Lien",
    "Senior Secured",
    "2030-09-30",
    "2031-06-30",
    "2029-03-31",
    "2028-12-31",
    "2032-07-15",
    # Compliance_Monitor
    "C-01",
    "C-04",
    "C-05",
    "C-10",
    "SECTOR",
    "Single Name",
    "Rating - CCC Bucket",
    "Rating - Unrated",
    "Instrument",
    "Sector",
    "Max single issuer exposure",
    "Max CCC+ and below (incl. Unrated and Other)",
    "Max unrated exposure",
    "Max 2nd Lien / Unsecured",
    "Max single sector (candidate's sector: Leisure Facilities & Services)",
    "Hard",
    "Soft",
    "issuer",
    "rating_bucket",
    "seniority",
    "sector",
    "Indenture §7.11 (synthetic)",
    "Watch",
    "analyst",
)


def test_every_template_text_cell_is_on_the_allow_list() -> None:
    book = openpyxl.load_workbook(TEMPLATE, read_only=True)
    found = {
        value
        for sheet in book.worksheets
        for row in sheet.iter_rows(values_only=True)
        for value in row
        if isinstance(value, str)
    }
    headers = {column for header in BRIEF_SHEETS.values() for column in header}
    allowed = set(TEMPLATE_TEXT) | headers | set(PORTFOLIO_README)
    assert sorted(found - allowed) == []  # a string nobody reviewed
    assert sorted(allowed - found) == []  # a stale entry


def test_readme_documents_every_sheet_column_and_fix() -> None:
    readme = (ROOT / "docs" / "templates" / "README.md").read_text(encoding="utf-8")
    section = readme[readme.index("## Portfolio workbook") :]
    assert all(f"`{sheet.name}`" in section for sheet in SHEETS)
    assert all(f"`{column.name}`" in section for sheet in SHEETS for column in sheet.columns)
    assert [fix for fix in FIXES if fix not in section] == []
    assert "deploy-vii validate-input portfolio" in section
    assert "sheet:Compliance_Monitor" in section
    assert "OK: 8 sheets, 5 compliance rows" in section


# Workbook structure: sheets, order and header pins.

SHEETS_FIX = (
    "delete the sheet or move it to another workbook; keep only Mandate, Exposure_Summary, "
    "Rating_Distribution, Sector_Exposure, Top_Holdings, Maturity_Profile, Compliance_Monitor, "
    "Holdings_Detail (optional) and README (optional)"
)
ORDER_FIX = (
    f"order the sheets as in {TEMPLATE_PATH}: Mandate, Exposure_Summary, Rating_Distribution, "
    "Sector_Exposure, Top_Holdings, Maturity_Profile, Compliance_Monitor, Holdings_Detail "
    "(optional) and README (optional)"
)


def test_a_ninth_unknown_sheet_is_refused() -> None:
    data = valid() | {"Holdings_Detail": [HOLDING], "Source Map": [{"Control": "Status"}]}
    report = check(data, headers={"Source Map": ("Control",)})
    assert rendered(report) == [
        f"{FILE}:Source Map!A1: unknown sheet 'Source Map' — fix: {SHEETS_FIX}"
    ]
    assert report.sheets == 8


def test_a_readme_sheet_is_accepted_anywhere() -> None:
    data = {"README": [{"note": "synthetic"}]} | valid()
    assert rendered(check(data, headers={"README": ("note",)})) == []


def test_a_missing_optional_sheet_is_accepted_and_a_missing_required_one_is_not() -> None:
    assert rendered(check(valid())) == []
    data = valid()
    del data["Top_Holdings"]
    found = "Mandate, Exposure_Summary, Rating_Distribution, Sector_Exposure, Maturity_Profile, "
    found += "Compliance_Monitor"
    assert rendered(check(data)) == [
        f"{FILE}:Top_Holdings!A1: no sheet named Top_Holdings (found: {found}) — fix: add the "
        f"sheet Top_Holdings with its header row, as in {TEMPLATE_PATH}"
    ]


def test_sheets_out_of_order_are_one_issue() -> None:
    data = valid()
    swapped = {"Exposure_Summary": data.pop("Exposure_Summary")} | data
    assert rendered(check(swapped)) == [
        f"{FILE}:Mandate!A1: sheet Mandate comes after Exposure_Summary — fix: {ORDER_FIX}"
    ]


@pytest.mark.parametrize(
    "second, refusal",
    [("cells", "Duplicate extracted locators"), ("empty", "Duplicate extraction coverage")],
)
def test_a_repeated_sheet_name_is_prepares_refusal(second: str, refusal: str) -> None:
    sheets = [
        (sheet.name, table_sheet(sheet.names, (), PORTFOLIO_ROWS[sheet.name], {}))
        for sheet in SHEETS
    ]
    empty = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>'
    repeated = sheets[0] if second == "cells" else ("Mandate", empty)
    (issue,) = validate_portfolio_bytes(FILE, package([sheets[0], repeated, *sheets[1:]])).issues
    assert (issue.sheet, issue.cell, issue.fix) == ("Mandate", "A1", NOT_OOXML_FIX)
    assert issue.problem.startswith("prepare cannot read this file as an OOXML workbook")
    assert refusal in issue.problem


def test_an_empty_sheet_and_a_header_only_sheet_are_one_issue_each() -> None:
    data = valid() | {"Maturity_Profile": []}
    report = check(data, headers={"Maturity_Profile": ()})
    assert rendered(report) == [
        f"{FILE}:Maturity_Profile!A1: sheet Maturity_Profile is empty — fix: copy the header "
        f"row of Maturity_Profile from {TEMPLATE_PATH} and add one row per maturity bucket"
    ]
    assert rendered(check(valid() | {"Compliance_Monitor": []})) == [
        f"{FILE}:Compliance_Monitor!A2: no rows below the header — fix: add one row per "
        "constraint, starting in row 2"
    ]


def header_fix(name: str) -> str:
    header = BRIEF_SHEETS[name]
    last = "ABCDEFGHIJKLM"[len(header) - 1]
    return f"set row 1 to exactly: {', '.join(header)} (A1 to {last}1, nothing after {last}1)"


@pytest.mark.parametrize("name", list(BRIEF_SHEETS))
def test_each_sheets_header_is_pinned(name: str) -> None:
    header = BRIEF_SHEETS[name]
    last = "ABCDEFGHIJKLM"[len(header) - 1]
    base = valid() | {"Holdings_Detail": [HOLDING]}
    swapped = (header[1], header[0], *header[2:])
    cases = [
        (swapped, "A1", f"header is '{header[1]}', expected '{header[0]}'"),
        (header[:-1], f"{last}1", f"header is empty, expected '{header[-1]}'"),
        (
            (*header, "extra"),
            "ABCDEFGHIJKLMN"[len(header)] + "1",
            f"header is 'extra', expected nothing after {last}1",
        ),
    ]
    for pinned, cell, problem in cases:
        report = check(base, headers={name: pinned})
        assert rendered(report) == [f"{FILE}:{name}!{cell}: {problem} — fix: {header_fix(name)}"]


# Hostile inputs: one issue each, never a traceback.

NOT_OOXML_FIX = "save the file as an Excel Workbook (.xlsx) and validate again"
DAMAGED_FIX = (
    "save the file again as an Excel Workbook (.xlsx) and validate again; prepare would also "
    "refuse this file"
)


def zip_with_workbook_member(payload: bytes, *, deflated: bool, encrypted: bool) -> bytes:
    """A one-member ZIP whose `xl/workbook.xml` header is patched to claim deflate
    compression or ZipCrypto encryption that its bytes do not honour (as in RV-1's tests)."""
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_STORED) as archive:
        archive.writestr("xl/workbook.xml", payload)
    raw = bytearray(stream.getvalue())
    central = raw.index(b"PK\x01\x02")
    if deflated:
        raw[8:10] = raw[central + 10 : central + 12] = ZIP_DEFLATED.to_bytes(2, "little")
    if encrypted:
        raw[6] |= 1
        raw[central + 8] |= 1
    return bytes(raw)


def zip_of(name: str, payload: str) -> bytes:
    stream = io.BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr(name, payload)
    return stream.getvalue()


@pytest.mark.parametrize(
    "raw, problem, fix",
    [
        (
            b"section,field\nStrategy,fund_name\n",
            "Preparation failed (BadZipFile): File is not a zip file",
            NOT_OOXML_FIX,
        ),
        (
            zip_of("notes.txt", "synthetic, not a workbook"),
            "Preparation failed (KeyError): \"There is no item named 'xl/workbook.xml' in the "
            'archive"',
            NOT_OOXML_FIX,
        ),
        (
            zip_with_workbook_member(b"\xff" * 64, deflated=True, encrypted=False),
            "damaged ZIP member (zlib.error: Error -3 while decompressing data: invalid block "
            "type)",
            DAMAGED_FIX,
        ),
        (
            zip_with_workbook_member(b"<workbook/>", deflated=False, encrypted=True),
            "damaged ZIP member (RuntimeError: File 'xl/workbook.xml' is encrypted, password "
            "required for extraction)",
            DAMAGED_FIX,
        ),
    ],
    ids=["not-a-zip", "zip-not-ooxml", "corrupt-deflate", "encrypted-member"],
)
def test_a_file_prepare_cannot_read_is_one_issue_not_a_traceback(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    raw: bytes,
    problem: str,
    fix: str,
) -> None:
    path = tmp_path / FILE
    path.write_bytes(raw)
    expected = (
        f"{FILE}:Mandate!A1: prepare cannot read this file as an OOXML workbook: {problem}"
        f" — fix: {fix}"
    )
    assert rendered(validate_portfolio(path)) == [expected]
    assert run_cli(["portfolio", str(path)], monkeypatch) == 1
    assert capsys.readouterr().out.splitlines() == [expected, "FAILED: 1 issues"]


def test_portfolio_facts_refuses_an_unreadable_source() -> None:
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(source(b"not a workbook"))
    assert [issue.fix for issue in refused.value.report.issues] == [NOT_OOXML_FIX]


def test_a_file_too_large_to_prepare_is_one_issue(monkeypatch: pytest.MonkeyPatch) -> None:
    from deploy_vii import preparation

    monkeypatch.setattr(preparation, "MAX_BYTES", 4)
    assert rendered(validate_portfolio(TEMPLATE)) == [
        "portfolio.template.xlsx:Mandate!A1: prepare cannot read this file as an OOXML"
        " workbook: Preparation failed (ValueError): Source exceeds extraction byte limit"
        f" — fix: {NOT_OOXML_FIX}"
    ]


def test_a_formula_in_a_header_or_label_cell_is_refused() -> None:
    report = check(changed("Rating_Distribution", 0, {"rating_bucket": '="B"'}))
    assert rendered(report) == [
        f"{FILE}:Rating_Distribution!A2: rating_bucket holds a formula; prepare freezes stored "
        f"values and never recalculates — fix: {FORMULA_FIX}"
    ]


def test_a_long_or_multiline_value_is_echoed_to_80_characters_on_one_line() -> None:
    long = "X" * 32000
    (issue,) = check(changed("Compliance_Monitor", 0, {"scope": long})).issues
    assert issue.problem.startswith(f"scope '{'X' * 80}…' is not one of")
    # The 80-character echo plus the twelve scopes in the problem and the fix.
    assert len(issue.render()) < 500
    (issue,) = check(changed("Compliance_Monitor", 0, {"breach_type": "Ha\nrd"})).issues
    assert "\n" not in issue.render()
    assert "breach_type 'Ha\\nrd' is not Hard or Soft" in issue.render()


# `prepare` round trip: the pack freezes every sheet under its `sheet:<name>` locators and the
# typed reader's locators name the frozen cells.


def test_prepare_freezes_the_template_with_a_locator_for_every_sheet(tmp_path: Path) -> None:
    shutil.copy(TEMPLATE, tmp_path / "portfolio.xlsx")
    run = Run(
        run_id="portfolio-template",
        issuer_id="example-holdings",
        issuer_name="Example Holdings Ltd (synthetic)",
        as_of=date(2026, 5, 29),
        case_kind="synthetic",
    )
    request = PreparationRequest(
        run=run,
        sources=(
            SourceRequest(source_id="portfolio", path="portfolio.xlsx", media_type=MEDIA_TYPE),
        ),
    )
    pack = prepare(request, tmp_path, tmp_path / "pack")
    assert load_pack(tmp_path / "pack") == pack
    (frozen,) = pack.sources
    assert frozen.status == "ready"
    spans = {span.locator: span for span in frozen.spans}
    for name, header in BRIEF_SHEETS.items():
        letters = "ABCDEFGHIJKLM"[: len(header)]
        assert [spans[f"sheet:{name}!{letter}1"].text for letter in letters] == list(header)
        stored = sum(1 for row in PORTFOLIO_ROWS[name] for value in row.values() if value)
        assert sum(1 for locator in spans if locator.startswith(f"sheet:{name}!")) == (
            len(header) + stored
        )
    assert not any(span.hidden or span.formula for span in frozen.spans)
    facts = portfolio_facts(frozen)
    assert facts == portfolio_facts(source(TEMPLATE.read_bytes(), TEMPLATE.name))
    assert spans[facts.mandate.locator("nav")].text == "630928502"
    for row in facts.compliance:
        assert Decimal(spans[row.locator("limit")].text) == row.limit
        assert Decimal(spans[row.locator("current_value")].text) == row.current_value
        assert spans[row.locator("id")].text == row.id
    for holding in facts.top_holdings:
        assert spans[holding.locator("issuer")].text == holding.issuer
    assert facts.holdings is not None
    for position in facts.holdings:
        assert spans[position.locator("maturity")].text == position.maturity.isoformat()


# Properties (decision 0012): generated valid workbooks pass through two independent writers;
# blanking one required cell is exactly one issue naming it; perturbing one headroom is exactly
# one arithmetic issue.

WORD = st.text(
    alphabet="abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789 .&()'-",
    min_size=1,
    max_size=20,
).filter(lambda value: value.strip() and "total" not in value.casefold())


def amount(low: str, high: str) -> st.SearchStrategy[Decimal]:
    return st.decimals(
        min_value=Decimal(low),
        max_value=Decimal(high),
        places=2,
        allow_nan=False,
        allow_infinity=False,
    )


def status_for(limit: Decimal, current: Decimal, watch: bool) -> str:
    if current > limit:
        return "Breach"
    return "Watch" if watch and limit - current < limit * Decimal("0.1") else "Pass"


Rows = dict[str, list[dict[str, str]]]


@st.composite
def portfolios(draw: st.DrawFn) -> Rows:
    """Valid synthetic portfolios: every rule's arithmetic and statuses consistent."""
    as_of = draw(st.dates(date(2000, 1, 1), date(2099, 12, 31))).isoformat()
    nav = str(draw(amount("1", "1000000000")))
    currency = draw(st.from_regex(r"[A-Z]{3}", fullmatch=True))
    mandate_rows = [
        {"section": "Strategy", "field": "fund_name", "value": "Fund " + draw(WORD)},
        {
            "section": "Strategy",
            "field": "strategy",
            "value": draw(st.sampled_from(("CLO", "Indexed Loans", "Indexed HY", "Other"))),
        },
        {"section": "Capital Structure", "field": "nav", "value": nav},
        {"section": "Capital Structure", "field": "nav_currency", "value": currency},
        {"section": "Reporting", "field": "as_of", "value": as_of},
    ]
    for index in range(draw(st.integers(0, 3))):
        mandate_rows.append(
            {
                "section": draw(st.sampled_from(SECTIONS.replace(" or ", ", ").split(", "))),
                "field": f"extra_{index}",
                "value": draw(WORD),
                "notes": draw(st.one_of(st.just(""), WORD)),
            }
        )
    exposure = [{"metric": "nav", "value": nav, "unit": currency, "as_of": as_of}]
    for index in range(draw(st.integers(0, 3))):
        exposure.append(
            {
                "metric": f"metric_{index}",
                "value": str(draw(amount("-1000", "1000"))),
                "unit": draw(WORD),
                "as_of": as_of,
            }
        )
    shares = [draw(amount("0", "20")) for _ in range(draw(st.integers(0, 4)))]
    shares.append(Decimal(100) - sum(shares, Decimal(0)) + draw(amount("-1", "1")))
    ratings = [
        {"rating_bucket": f"R{index}", "pct_nav": str(share), "issuers": str(index)}
        for index, share in enumerate(shares)
    ]
    sectors = []
    for index in range(draw(st.integers(1, 4))):
        share = draw(amount("0", "20"))
        row = {"sector": f"Sector {index}", "pct_nav": str(share), "issuers": str(index)}
        if draw(st.booleans()):
            ceiling = draw(amount("0", "30"))
            row["limit_pct_nav"] = str(ceiling)
            if draw(st.booleans()):
                row["headroom_pct_nav"] = str(ceiling - share)
            if draw(st.booleans()):
                row["status"] = status_for(ceiling, share, draw(st.booleans()))
        sectors.append(row)
    top = [
        {
            "rank": str(rank),
            "issuer": "Issuer " + draw(WORD),
            "instrument": draw(WORD),
            "pct_nav": str(draw(amount("0", "5"))),
            "rating": draw(WORD),
        }
        for rank in range(1, draw(st.integers(1, 5)) + 1)
    ]
    maturities = [
        {"bucket": f"Bucket {index}", "pct_nav": str(draw(amount("0", "60")))}
        for index in range(draw(st.integers(1, 4)))
    ]
    compliance = []
    for index in range(draw(st.integers(1, 5))):
        ceiling, current = draw(amount("0", "50")), draw(amount("0", "60"))
        compliance.append(
            {
                "id": f"C-{index:02d}" if index else "SECTOR",
                "constraint_category": draw(WORD),
                "parameter": draw(WORD),
                "limit": str(ceiling),
                "limit_unit": "% NAV" if index == 0 or draw(st.booleans()) else currency,
                "breach_type": draw(st.sampled_from(("Hard", "Soft"))),
                "scope": draw(st.sampled_from(("issuer", "sector", "rating_bucket", "other"))),
                "source_document": draw(WORD),
                "current_value": str(current),
                "headroom": str(ceiling - current),
                "status": status_for(ceiling, current, draw(st.booleans())),
                "last_checked": as_of,
                "checked_by": "analyst",
            }
        )
    data: Rows = {
        "Mandate": mandate_rows,
        "Exposure_Summary": exposure,
        "Rating_Distribution": ratings,
        "Sector_Exposure": sectors,
        "Top_Holdings": top,
        "Maturity_Profile": maturities,
        "Compliance_Monitor": compliance,
    }
    if draw(st.booleans()):
        data["Holdings_Detail"] = [
            HOLDING | {"facility_id": f"F{index}", "pct_nav": str(draw(amount("0", "5")))}
            for index in range(draw(st.integers(1, 3)))
        ]
    return data


@settings(max_examples=40, deadline=None)
@given(portfolios())
def test_valid_generated_workbooks_are_ok_through_two_writers(data: Rows) -> None:
    report = validate_portfolio_bytes(FILE, portfolio_bytes(data))
    assert rendered(report) == []
    assert report.sheets == len(data)
    assert report.compliance_rows == len(data["Compliance_Monitor"])
    # Differential: an independent writer (openpyxl, every value a text cell) reads the same.
    assert validate_portfolio_bytes(FILE, workbook(dict(data))) == report


def required_cells(data: Rows) -> list[tuple[str, int, str]]:
    """Every required cell whose blanking can only be its own issue: the section and field of
    the five required Mandate rows and the metric of the nav row identify a required row, so
    blanking them also reports the row missing (tested separately)."""
    optional = {name: OPTIONAL_COLUMNS.get(name, set()) for name in BRIEF_SHEETS}
    return [
        (name, index, column)
        for name, rows in data.items()
        for index, row in enumerate(rows)
        for column in BRIEF_SHEETS[name]
        if column not in optional[name]
        and not (name == "Mandate" and index < 5 and column in ("section", "field"))
        and not (name == "Exposure_Summary" and index == 0 and column == "metric")
    ]


@settings(max_examples=60, deadline=None)
@given(portfolios(), st.data())
def test_blanking_one_required_cell_is_exactly_one_issue_naming_it(
    data: Rows, choice: st.DataObject
) -> None:
    name, index, column = choice.draw(st.sampled_from(required_cells(data)))
    data[name][index] = data[name][index] | {column: ""}
    report = validate_portfolio_bytes(FILE, portfolio_bytes(data))
    letter = "ABCDEFGHIJKLM"[BRIEF_SHEETS[name].index(column)]
    assert [(issue.sheet, issue.cell, issue.problem) for issue in report.issues] == [
        (name, f"{letter}{index + 2}", f"{column} is empty")
    ]


@settings(max_examples=60, deadline=None)
@given(portfolios(), st.data())
def test_perturbing_one_headroom_is_exactly_one_arithmetic_issue(
    data: Rows, choice: st.DataObject
) -> None:
    targets = [
        ("Compliance_Monitor", index, "headroom", "J")
        for index, row in enumerate(data["Compliance_Monitor"])
        if row["limit_unit"] == "% NAV"
    ] + [
        ("Sector_Exposure", index, "headroom_pct_nav", "E")
        for index, row in enumerate(data["Sector_Exposure"])
        if row.get("headroom_pct_nav")
    ]
    name, index, column, letter = choice.draw(st.sampled_from(targets))
    shift = choice.draw(st.one_of(amount("0.02", "100"), amount("-100", "-0.02")))
    row = data[name][index]
    data[name][index] = row | {column: str(Decimal(row[column]) + shift)}
    report = validate_portfolio_bytes(FILE, portfolio_bytes(data))
    assert [(issue.sheet, issue.cell) for issue in report.issues] == [
        (name, f"{letter}{index + 2}")
    ]
    assert "(to within 0.01)" in report.issues[0].problem


# Fix round 1 (review F1): a number Excel shows as a percentage stores the fraction (2.5 % is
# stored 0.025), so in a percent column it would be read 100 times too small.

PERCENT_FIX = "format the cell as General and enter 2.5 for 2.5 % NAV"


def formatted(data: dict[str, list[dict[str, Any]]], formats: dict[tuple[str, str], str]) -> bytes:
    """`workbook(data)` with number formats set on (sheet, cell) through openpyxl."""
    book = openpyxl.load_workbook(io.BytesIO(workbook(data)))
    for (sheet, cell), number_format in formats.items():
        book[sheet][cell].number_format = number_format
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_percent_formatted_limits_on_a_pct_nav_row_are_refused() -> None:
    data = changed(
        "Compliance_Monitor", 0, {"limit": 0.025, "current_value": 0.0237, "headroom": 0.0013}
    )
    cells = {("Compliance_Monitor", cell): "0.00%" for cell in ("D2", "I2", "J2")}
    report = validate_portfolio_bytes(FILE, formatted(data, cells))
    assert rendered(report) == [
        f"{FILE}:Compliance_Monitor!D2: limit '0.025' is formatted as a percentage — fix: "
        f"{PERCENT_FIX}",
        f"{FILE}:Compliance_Monitor!I2: current_value '0.0237' is formatted as a percentage"
        f" — fix: {PERCENT_FIX}",
        f"{FILE}:Compliance_Monitor!J2: headroom '0.0013' is formatted as a percentage — fix: "
        f"{PERCENT_FIX}",
    ]


@pytest.mark.parametrize("number_format", ["0%", "0.00%", "0.000%", "#,##0.0%;-0.0%"])
def test_a_percent_formatted_pct_nav_is_one_issue(number_format: str) -> None:
    data = changed("Rating_Distribution", 0, {"pct_nav": 0.605})
    raw = formatted(data, {("Rating_Distribution", "B2"): number_format})
    assert rendered(validate_portfolio_bytes(FILE, raw)) == [
        f"{FILE}:Rating_Distribution!B2: pct_nav '0.605' is formatted as a percentage — fix: "
        f"{PERCENT_FIX}"
    ]


@pytest.mark.parametrize("number_format", ['0.00"%"', "0.00\\%", "0.00", "General"])
def test_a_literal_percent_sign_or_a_plain_format_is_accepted(number_format: str) -> None:
    data = changed("Rating_Distribution", 0, {"pct_nav": 60.5})
    raw = formatted(data, {("Rating_Distribution", "B2"): number_format})
    assert rendered(validate_portfolio_bytes(FILE, raw)) == []


def test_percent_format_is_judged_only_in_percent_columns() -> None:
    # An amount limit in the NAV currency and a count are not percentages of NAV.
    row = {"limit": 20000000, "limit_unit": "USD", "current_value": 15000000, "headroom": 1}
    data = changed("Compliance_Monitor", 0, row | {"status": "Pass"})
    cells = {("Compliance_Monitor", "D2"): "0%", ("Rating_Distribution", "C2"): "0%"}
    assert rendered(validate_portfolio_bytes(FILE, formatted(data, cells))) == []


# Fix round 1 (review F2): Excel error cells (`t="e"`) and TRUE/FALSE cells (`t="b"`, stored
# 1 or 0) are not values, in any column.

ERROR_FIX = "enter the real value in place of the error, or leave an optional cell empty"
BOOLEAN_FIX = "enter text or a plain number, not TRUE or FALSE"


@pytest.mark.parametrize(
    "sheet, index, change, cell, problem, fix",
    [
        (
            "Top_Holdings",
            0,
            {"issuer": "#N/A"},
            "B2",
            "issuer holds the Excel error '#N/A'",
            ERROR_FIX,
        ),
        (
            "Compliance_Monitor",
            0,
            {"current_value": "#DIV/0!"},
            "I2",
            "current_value holds the Excel error '#DIV/0!'",
            ERROR_FIX,
        ),
        ("Mandate", 5, {"notes": "#REF!"}, "D7", "notes holds the Excel error '#REF!'", ERROR_FIX),
        (
            "Rating_Distribution",
            0,
            {"pct_nav": True},
            "B2",
            "pct_nav holds a TRUE/FALSE cell",
            BOOLEAN_FIX,
        ),
        ("Top_Holdings", 1, {"rating": False}, "E3", "rating holds a TRUE/FALSE cell", BOOLEAN_FIX),
    ],
)
def test_error_and_boolean_cells_are_one_issue_each(
    sheet: str, index: int, change: dict[str, Any], cell: str, problem: str, fix: str
) -> None:
    raw = workbook(changed(sheet, index, change))
    kinds = {span.data_type for span in source(raw).spans}
    assert kinds >= {"e"} or kinds >= {"b"}  # openpyxl wrote a real error or boolean cell
    assert rendered(validate_portfolio_bytes(FILE, raw)) == [
        f"{FILE}:{sheet}!{cell}: {problem} — fix: {fix}"
    ]


# Fix round 2: projection must never change the frozen cell identity.
def appended_cells(addresses: tuple[tuple[str, str], ...]) -> bytes:
    raw = workbook(valid())
    stream = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as original, ZipFile(stream, "w") as rewritten:
        for member in original.infolist():
            content = original.read(member)
            if member.filename == "xl/worksheets/sheet7.xml":
                cells = "".join(
                    f"<c r={quoteattr(address)}><v>{value}</v></c>" for address, value in addresses
                )
                content = content.replace(
                    b"</sheetData>", f'<row r="2">{cells}</row></sheetData>'.encode()
                )
            rewritten.writestr(member, content)
    return stream.getvalue()


def test_newline_cell_aliases_cannot_replace_frozen_financial_cells() -> None:
    raw = appended_cells((("D2\n", "250"), ("I2\n", "237"), ("J2\n", "13")))
    frozen = source(raw)
    spans = {span.locator: span.text for span in frozen.spans}
    assert spans["sheet:Compliance_Monitor!D2"] == "2.5"
    assert spans["sheet:Compliance_Monitor!D2\n"] == "250"
    result = validate_portfolio_bytes(FILE, raw)
    assert not result.ok
    assert "canonical" in result.issues[0].problem
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(frozen)
    assert refused.value.report == result


@pytest.mark.parametrize(
    "address",
    ["D" + "1" * 4301, "A" * 4301 + "2", "D1048577", "XFE2", "D02", "D0", "d2", " D2", "D2 "],
    ids=[
        "huge-row",
        "huge-column",
        "row-limit",
        "column-limit",
        "leading-zero",
        "zero",
        "lower",
        "leading-space",
        "trailing-space",
    ],
)
def test_noncanonical_coordinates_return_a_report_not_a_traceback(address: str) -> None:
    raw = appended_cells(((address, "250"),))
    result = validate_portfolio_bytes(FILE, raw)
    assert not result.ok
    assert "canonical" in result.issues[0].problem
    with pytest.raises(PortfolioRefused):
        portfolio_facts(source(raw))


@pytest.mark.parametrize("address", ["A1", "XFD1048576"])
def test_native_boundary_cells_have_their_exact_frozen_identity(address: str) -> None:
    # A1 already exists, so use an uncontracted README sheet for the first native cell.
    frozen = source(workbook(valid()))
    payload = frozen.model_dump()
    payload["spans"] += ({"locator": f"sheet:README!{address}", "text": "ignored"},)
    payload["coverage"] += ({"locator": "sheet:README", "status": "retained"},)
    assert portfolio_facts(Source.model_validate(payload)).compliance[0].limit == Decimal("2.5")
    if address.startswith("XFD"):
        assert validate_portfolio_bytes(FILE, appended_cells(((address, "250"),))).ok


@pytest.mark.parametrize("target", ["spans", "coverage", "both", "collision"])
def test_noncanonical_sheet_uris_cannot_alias_frozen_cells(target: str) -> None:
    payload = source(workbook(valid())).model_dump()
    alias = "sheet:%43ompliance_Monitor"
    if target in ("spans", "both"):
        payload["spans"] = tuple(
            span | {"locator": span["locator"].replace("sheet:Compliance_Monitor", alias)}
            for span in payload["spans"]
        )
    if target in ("coverage", "both"):
        payload["coverage"] = tuple(
            item | {"locator": item["locator"].replace("sheet:Compliance_Monitor", alias)}
            for item in payload["coverage"]
        )
    if target == "collision":
        payload["spans"] += ({"locator": alias + "!D2", "text": "250"},)
        payload["coverage"] += ({"locator": alias, "status": "retained"},)
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(Source.model_validate(payload))
    assert "canonical" in refused.value.report.issues[0].problem


@pytest.mark.parametrize("status", ["missing", "unreadable"])
def test_unavailable_prepared_source_is_refused(status: str) -> None:
    payload = source(workbook(valid())).model_dump()
    payload.update(status=status, extraction_status="blocked", limitation="synthetic failure")
    if status == "missing":
        payload.update(sha256=None, spans=())
    frozen = Source.model_validate(payload)
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(frozen)
    assert "cannot read" in refused.value.report.issues[0].problem


@pytest.mark.parametrize("name", list(BRIEF_SHEETS))
def test_unreadable_contracted_sheet_coverage_is_refused(name: str) -> None:
    payload = source(workbook(valid() | {"Holdings_Detail": [HOLDING]})).model_dump()
    payload["coverage"] = tuple(
        item
        | (
            {"status": "unreadable", "limitation": "synthetic failure"}
            if item["locator"] == f"sheet:{name}"
            else {}
        )
        for item in payload["coverage"]
    )
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(Source.model_validate(payload))
    assert refused.value.report.issues[0].sheet == name
    assert "unreadable" in refused.value.report.issues[0].problem


def test_empty_coverage_cannot_supply_frozen_data_rows() -> None:
    payload = source(workbook(valid())).model_dump()
    payload["coverage"] = tuple(
        item | ({"status": "empty"} if item["locator"] == "sheet:Compliance_Monitor" else {})
        for item in payload["coverage"]
    )
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(Source.model_validate(payload))
    assert "sheet Compliance_Monitor is empty" in refused.value.report.issues[0].problem


def test_degraded_source_and_unreadable_optional_readme_remain_valid() -> None:
    payload = source(workbook(valid())).model_dump()
    assert payload["status"] == "ready" and payload["extraction_status"] == "degraded"
    payload["coverage"] += ({"locator": "sheet:README", "status": "unreadable"},)
    assert portfolio_facts(Source.model_validate(payload)).holdings is None
    raw = workbook(valid() | {"Holdings_Detail": []}, headers={"Holdings_Detail": ()})
    with pytest.raises(PortfolioRefused) as refused:
        portfolio_facts(source(raw))
    assert "sheet Holdings_Detail is empty" in refused.value.report.issues[0].problem


def test_cli_refuses_noncanonical_cells_with_exit_one(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / FILE
    path.write_bytes(appended_cells((("D2\n", "250"),)))
    assert run_cli(["portfolio", str(path)], monkeypatch) == 1
    assert "FAILED:" in capsys.readouterr().out


@settings(max_examples=40, deadline=None)
@given(portfolios())
def test_generated_typed_fields_resolve_to_the_actual_frozen_cells(data: Rows) -> None:
    frozen = source(portfolio_bytes(data))
    facts = portfolio_facts(frozen)
    for name in ("fund_name", "strategy", "nav", "nav_currency", "as_of"):
        (span,) = (span for span in frozen.spans if span.locator == facts.mandate.locator(name))
        value = getattr(facts.mandate, name)
        if isinstance(value, Decimal):
            assert value == Decimal(span.text.strip())
        elif isinstance(value, date):
            assert value == date.fromisoformat(span.text.strip())
        else:
            assert value == span.text.strip()
    groups = (
        facts.mandate.fields,
        facts.exposures,
        facts.ratings,
        facts.sectors,
        facts.top_holdings,
        facts.maturities,
        facts.compliance,
        facts.holdings or (),
    )
    for rows in groups:
        for row in rows:
            expected_row = data[row.SHEET][row.row - 2]
            for index, column in enumerate(BRIEF_SHEETS[row.SHEET], 1):
                locator = f"sheet:{row.SHEET}!{'ABCDEFGHIJKLM'[index - 1]}{row.row}"
                assert row.locator(column) == locator
                matches = [span for span in frozen.spans if span.locator == locator]
                expected = expected_row.get(column, "").strip()
                value = getattr(row, column)
                if not expected:
                    assert value is None
                    continue
                assert len(matches) == 1 and matches[0].text.strip() == expected
                actual = matches[0].text.strip()
                if isinstance(value, Decimal):
                    assert value == Decimal(actual)
                elif isinstance(value, int):
                    assert value == int(actual)
                elif isinstance(value, date):
                    assert value == date.fromisoformat(actual)
                else:
                    assert value == actual


@settings(max_examples=40, deadline=None)
@given(portfolios(), st.sampled_from(["\n", " ", "\r", "1" * 4301, "!D2"]))
def test_generated_aliases_cannot_overwrite_typed_frozen_values(data: Rows, suffix: str) -> None:
    payload = source(portfolio_bytes(data)).model_dump()
    payload["spans"] += ({"locator": "sheet:Compliance_Monitor!D2" + suffix, "text": "999"},)
    with pytest.raises(PortfolioRefused):
        portfolio_facts(Source.model_validate(payload))
